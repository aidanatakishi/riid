# -*- coding: utf-8 -*-
"""PPTX hesabatını oxuyur və qurumun rəqəmsal inkişafını qiymətləndirir."""
from __future__ import annotations

import os
import re
import zipfile
from collections import Counter
from xml.etree import ElementTree as ET

DIR_NEEDLES = (
    ('strategiya', ('strategiya', 'strategy', 'strategic', 'yol xeritesi', 'kpi')),
    ('xidmetler', ('xidmet', 'service', 'portal', 'e xidmet', 'onlayn')),
    ('texniki', ('texniki', 'infrastruktur', 'texnoloj', 'technical', 'bulud', 'g cloud', 'kiber', 'data')),
    ('emeliyyat', ('emeliyyat', 'meliyyat', 'operation', 'proses', 'agile', 'teskilati')),
)

DIR_TITLES = {
    'strategiya': 'Strategiya',
    'xidmetler': 'Xidmətlər',
    'texniki': 'Texniki-texnoloji infrastruktur',
    'emeliyyat': 'Əməliyyat modelləri',
}

ISQ_PILLARS = (
    ('strategy', 'Strateji və idarəçilik sənədləri', 'Strateji',
     ('strateji ve idarecilik', 'idarecilik sened', 'strateji sened')),
    ('procedure', 'Prosedurlar və təlimatlar', 'Prosedur',
     ('prosedurlar ve telimat', 'prosedur ve telimat')),
    ('ops', 'Əməliyyatlar', 'Əməliyyat',
     ('emeliyyatlar',)),
    ('lifecycle', 'Proqram təminatının həyat dövrü', 'Həyat dövrü',
     ('hayat dovru', 'proqram teminatinin hayat')),
    ('quality', 'Proqram təminatının keyfiyyətinə nəzarət', 'Keyfiyyət',
     ('keyfiyyetine nezeret', 'proqram teminatinin keyfiyyet')),
    ('integ', 'İnteqrasiya', 'İnteqrasiya',
     ('inteqrasiya',)),
    ('info', 'İnformasiya ehtiyatı sistemi', 'Ehtiyat',
     ('informasiya ehtiyati', 'ehtiyati sistemi')),
    ('network', 'Şəbəkə', 'Şəbəkə',
     ('sebeke',)),
    ('data', 'Məlumatların idarəedilməsi', 'Məlumat',
     ('melumatlarin idareedilmesi', 'melumatlarin idare')),
    ('infra', 'Fiziki və virtual infrastruktur', 'İnfrastruktur',
     ('fiziki ve virtual', 'virtual infrastruktur')),
)

EXQ_STAR_WEIGHTS = (0.2, 0.4, 0.6, 0.8, 1.0)

THEMES = (
    ('strategiya', 'Strateji idarəetmə', ('strategiya', 'yol xeritesi', 'hedef', 'baxis', 'kpi', 'elaqelendir')),
    ('e_xidmet', 'Elektron xidmətlər', ('e xidmet', 'elektron xidmet', 'portal', 'onlayn', 'g2c', 'g2b')),
    ('data', 'Məlumat idarəetməsi', ('melumat', 'data', 'analitik', 'statistika', 'reyestr')),
    ('cloud', 'Bulud və infrastruktur', ('bulud', 'g cloud', 'infrastruktur', 'server', 'datacenter')),
    ('cyber', 'Kiber təhlükəsizlik', ('kiber', 'tehlukesizlik', 'backup', 'ehtiyat', 'risk')),
    ('ai', 'Süni intellekt', ('intellekt', 'ai', 'machine learning', 'avtomatlasdir')),
    ('hr', 'Rəqəmsal bacarıqlar', ('savad', 'telim', 'kadr', 'bacariq', 'insan resurs')),
    ('legal', 'Hüquqi tənzimləmə', ('huquqi', 'normativ', 'qanun', 'reyestr', 'e imza')),
)

NEG = (
    'catismaz', 'zeif', 'yoxdur', 'proble', 'risk', 'gerilik', 'qeyri kafi',
    'formalasdirilmayib', 'formalesdirilmeyib', 'tesdiq edilmeyib', 'qurulmayib',
    ' aparilmir', 'kagiz', 'el ile', 'manual', 'pərakəndə', 'perakende', 'asili',
    'heyata kecirilmir', 'izlenilmir', 'olculmur', 'bosluq', 'eksiklik',
)
POS = (
    'tetbiq olunub', 'tesdiq olunub', 'movcuddur', 'heyata kecirilir',
    'avtomat', 'inteqrasiya', 'tekmilles', 'merkezles', 'proaktiv',
    'real vaxtda', 'olculebilen', 'ugurla',
)
ACT = (
    'tovsiye', 'edilmelidir', 'planlasdirilsin', 'lazimdir', 'teleb olunur',
    'gorulmelidir', 'hazirlanmalidir', 'guclendirilsin',
)

_TR = str.maketrans({
    'ı': 'i', 'İ': 'i', 'I': 'i', 'ə': 'e', 'Ə': 'e',
    'ö': 'o', 'Ö': 'o', 'ü': 'u', 'Ü': 'u',
    'ğ': 'g', 'Ğ': 'g', 'ş': 's', 'Ş': 's', 'ç': 'c', 'Ç': 'c',
})

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}


def fold(value):
    s = '' if value is None else str(value)
    s = s.replace('\xa0', ' ').translate(_TR).lower()
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


def cell_text(value):
    if value is None:
        return ''
    s = str(value).replace('\xa0', ' ').strip()
    return '' if s.lower() in ('none', 'null', '-', '—') else s


def parse_num(value):
    if value is None or value == '':
        return None
    s = str(value).strip().replace(',', '.')
    m = re.search(r'-?\d+(?:\.\d+)?', s)
    if not m:
        return None
    n = float(m.group(0))
    return n if n == n and 0 <= n <= 100 else None


def sentences(text):
    parts = re.split(r'(?<=[.!?…])\s+|\n+', text or '')
    out = []
    for p in parts:
        t = re.sub(r'\s+', ' ', p).strip(' •\t-–—')
        if len(t) >= 18:
            out.append(t)
    return out


def dir_id_from_text(text):
    f = fold(text)
    if not f:
        return ''
    for did, needles in DIR_NEEDLES:
        if any(n in f for n in needles):
            if did == 'xidmetler' and 'say' in f and 'xidmet' not in f:
                continue
            return did
    return ''


def polarity(text):
    f = fold(text)
    neg = sum(1 for n in NEG if n in f)
    pos = sum(1 for n in POS if n in f)
    act = sum(1 for n in ACT if n in f)
    if act and not pos:
        return 'action'
    if neg > pos:
        return 'finding'
    if pos > neg:
        return 'strength'
    if act:
        return 'action'
    return ''


def year_from_text(text):
    years = [int(y) for y in re.findall(r'(20[2-3]\d)', text or '')]
    if not years:
        return None
    counts = Counter(years)
    return sorted(counts, key=lambda y: (counts[y], y), reverse=True)[0]


_ORG_HINTS = (
    ('dovlet vergi xidmet', 'Dövlət Vergi Xidməti'),
    ('vergi xidmeti', 'Dövlət Vergi Xidməti'),
)


def org_from_text(chunks):
    joined = '\n'.join(chunks[:8])
    fj = fold(joined)
    for needle, name in _ORG_HINTS:
        if needle in fj:
            return name
    m = re.search(
        r'(?i)(?:qurum|təşkilat|müəssisə|nazirlik|komitə|agentlik)\s*[:\-–]\s*([^\n|]{4,80})',
        joined,
    )
    if m:
        return clean_org(m.group(1))
    for chunk in chunks[:6]:
        name = clean_org(chunk)
        if looks_like_org(name):
            return name
    return ''


def clean_org(text):
    t = re.sub(r'\s+', ' ', cell_text(text))
    t = re.sub(r'(?i)(diaqnostika|hesabat|reqemsallasma|qiymetlendirme|prezentasiya).*$', '', t).strip(' .,-')
    return t[:80]


def looks_like_org(name):
    f = fold(name)
    if not f or len(f) < 6:
        return False
    return any(tok in f for tok in (
        'nazirlik', 'komite', 'agentlik', 'xidmeti', 'idaresi', 'muessise',
        'qurumu', 'bank', 'sirket', 'sehiyye', 'elm', 'tehsil', 'vergi',
    ))


def clean_slide_title(text):
    raw = cell_text(text).replace('\x0b', '\n').replace('\x0c', '\n')
    lines = [re.sub(r'\s+', ' ', ln).strip(' •\t-–—') for ln in raw.split('\n')]
    lines = [ln for ln in lines if ln]
    if not lines:
        return ''
    title = lines[0]
    i = 1
    while (
        i < len(lines)
        and len(title) < 120
        and len(lines[i]) <= 48
        and not re.search(r'[.!?]$', title)
    ):
        title = title + ' ' + lines[i]
        i += 1
    return title


def extract_pptx_slides(path):
    slides = []
    try:
        from pptx import Presentation
    except ImportError:
        return extract_pptx_zip(path)
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, start=1):
        bits = []
        title = ''
        try:
            ph = slide.shapes.title
            if ph is not None and getattr(ph, 'has_text_frame', False):
                title = clean_slide_title(ph.text_frame.text)
        except Exception:
            title = ''
        for shape in slide.shapes:
            if getattr(shape, 'has_text_frame', False):
                text = cell_text(shape.text_frame.text)
                if text:
                    bits.append(text)
                    if not title:
                        title = clean_slide_title(text)
            if getattr(shape, 'has_table', False):
                table = shape.table
                for row in table.rows:
                    cells = [cell_text(c.text) for c in row.cells]
                    line = ' | '.join(x for x in cells if x)
                    if line:
                        bits.append(line)
        notes = ''
        try:
            if slide.has_notes_slide:
                notes = cell_text(slide.notes_slide.notes_text_frame.text)
        except Exception:
            notes = ''
        if notes:
            bits.append(notes)
        if bits:
            joined = clean_slide_title('\n'.join(bits[:6]))
            if joined and (not title or len(title) < 32):
                if not re.search(r'\(\s*\d', joined) and len(joined) > len(title or ''):
                    title = joined
        body = '\n'.join(bits)
        slides.append({
            'n': i,
            'title': title or ('Slayd %s' % i),
            'text': body,
        })
    return slides


def extract_pptx_zip(path):
    slides = []
    with zipfile.ZipFile(path) as zf:
        names = sorted(
            n for n in zf.namelist()
            if re.match(r'ppt/slides/slide\d+\.xml$', n)
        )
        for i, name in enumerate(names, start=1):
            xml = zf.read(name)
            root = ET.fromstring(xml)
            texts = [cell_text(el.text) for el in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t')]
            texts = [t for t in texts if t]
            body = '\n'.join(texts)
            title = clean_slide_title(texts[0]) if texts else ('Slayd %s' % i)
            slides.append({'n': i, 'title': title, 'text': body})
    return slides


def extract_pdf_pages(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError('PDF oxumaq üçün pypdf paketi lazımdır.')
    try:
        reader = PdfReader(path)
    except Exception as e:
        raise RuntimeError('PDF oxunmadı: %s' % e)
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ''
        except Exception:
            raw = ''
        raw = raw.replace('\xa0', ' ')
        title = clean_slide_title(raw)
        pages.append({
            'n': i,
            'title': title or ('Səhifə %s' % i),
            'text': cell_text(raw),
        })
    return pages


_OFFICIAL_LABELS = {
    'strategiya': ('strategiya', 'strategiya istiqameti'),
    'xidmetler': ('xidmetler',),
    'texniki': (
        'texniki',
        'texniki texnoloji',
        'texniki texnoloji infrastruktur',
    ),
    'emeliyyat': (
        'emeliyyat',
        'emeliyyat modeli',
        'emeliyyat modelleri',
    ),
}

_OVERALL_LABELS = (
    'umumi reqemsal yetkinlik',
    'umumi yetkinlik',
    'yekun reqemsallasma',
    'yekun netice',
    'umumi netice',
    'movcud reqemsallasma seviyyesi',
    'reqemsal yetkinlik',
)

_CRITERIA = (
    ('e_xidmet', 'Rəqəmsal xidmətlər', ('reqemsal xidmetler',)),
    ('e_xidmet', 'Əlçatanlıq və istifadəçi təcrübəsi', ('elcatanliq', 'istifadeci tecrubesi')),
    ('e_xidmet', 'Rəqəmsal xidmətlərin təşkili', ('xidmetlerin teskili ve tekmillesdirilmesi',)),
    ('cloud', 'İnformasiya infrastrukturu', ('informasiya infrastrukturu',)),
    ('ai', 'Süni intellekt', ('suni intellekt hellerinin', 'suni intellekt')),
    ('data', 'Məlumat idarəetməsi', ('melumat idareetmesi',)),
    ('strategiya', 'Rəqəmsal strategiya və əlaqələndirmə', ('reqemsal strategiya ve elaqelendirme',)),
    ('emeliyyat', 'Təşkilati struktur', ('teskilati struktur',)),
    ('hr', 'Bilik və bacarıqların artırılması', ('bilik ve bacariq',)),
    ('legal', 'Hüquqi tənzimləmə', ('huquqi tenzimleme',)),
    ('texniki', 'Əməkdaşlıq və inteqrasiya', ('emekdasliq ve inteqrasiya',)),
    ('emeliyyat', 'Daxili proseslərin avtomatlaşdırılması', ('daxili idareetme proseslerinin',)),
)


def official_dir_label(label):
    f = fold(label)
    if not f:
        return ''
    if f in _OVERALL_LABELS:
        return 'overall'
    for did, names in _OFFICIAL_LABELS.items():
        if f in names:
            return did
    return ''


def standalone_num(line):
    s = cell_text(line).replace(',', '.')
    if re.fullmatch(r'\d{1,3}(?:\.\d+)?', s):
        return parse_num(s)
    return None


def split_label_score(line):
    raw = cell_text(line)
    if not raw or len(raw) > 90:
        return '', None
    m = re.match(
        r'^\s*(.+?)\s*[\(\[]\s*(\d{1,3}(?:[.,]\d+)?)\s*[\)\]]\s*$',
        raw,
    )
    if not m:
        m = re.match(r'^\s*(.+?)\s+(\d{1,3}(?:[.,]\d+)?)\s*$', raw)
    if not m:
        return fold(raw), None
    return fold(m.group(1)), parse_num(m.group(2))


def normalize_slide_lines(text):
    raw = (text or '').replace('\x0b', '\n').replace('\r', '\n')
    parts = [re.sub(r'\s+', ' ', p).strip(' •\t-–—') for p in raw.split('\n')]
    parts = [p for p in parts if p]
    merged = []
    for p in parts:
        if merged:
            prev_f = fold(merged[-1])
            cur_f = fold(p)
            if prev_f.endswith('texnoloji') and cur_f.startswith('infrastruktur'):
                merged[-1] = merged[-1] + ' ' + p
                continue
            if prev_f.endswith('istifadeci') and 'tecrube' in cur_f:
                merged[-1] = merged[-1] + ' ' + p
                continue
        merged.append(p)
    return merged


def pick_unique_score(votes):
    if not votes:
        return None
    counts = Counter(votes)
    return sorted(counts, key=lambda n: (counts[n], -abs(n - 50)), reverse=True)[0]


def apply_label_nums(store, labels, nums):
    if not nums:
        return None
    overall = None
    mapped = []
    if labels and len(nums) == len(labels) + 1:
        overall = nums[0]
        mapped = list(zip(labels, nums[1:]))
    elif labels and len(nums) == len(labels):
        mapped = list(zip(labels, nums))
    elif not labels and len(nums) == 1:
        overall = nums[0]
    for did, num in mapped:
        if did and num is not None:
            store.setdefault(did, []).append(num)
    return overall


def collect_scores(slides):
    current_votes = {d: [] for d in DIR_TITLES}
    target_votes = {d: [] for d in DIR_TITLES}
    overall_votes = []
    target_overall = []
    criteria = []

    for sl in slides or []:
        lines = normalize_slide_lines((sl.get('title') or '') + '\n' + (sl.get('text') or ''))
        i = 0
        while i < len(lines):
            line = lines[i]
            folded = fold(line)
            num = standalone_num(line)
            label_f, label_num = split_label_score(line)
            kind = official_dir_label(label_f if label_num is not None else folded)

            if folded in ('hedef olunan seviyye', 'hedef olunan'):
                labels, nums = [], []
                j = i + 1
                while j < len(lines):
                    did = official_dir_label(fold(lines[j]))
                    n = standalone_num(lines[j])
                    if did and did != 'overall':
                        labels.append(did)
                        j += 1
                        continue
                    if n is not None:
                        nums.append(n)
                        j += 1
                        continue
                    break
                y = apply_label_nums(target_votes, labels, nums)
                if y is not None:
                    target_overall.append(y)
                i = j
                continue

            if folded in ('movcud reqemsallasma seviyyesi', 'movcud seviyye'):
                labels, nums = [], []
                j = i + 1
                while j < len(lines):
                    did = official_dir_label(fold(lines[j]))
                    n = standalone_num(lines[j])
                    if did and did != 'overall':
                        labels.append(did)
                        j += 1
                        continue
                    if n is not None:
                        nums.append(n)
                        j += 1
                        continue
                    break
                if labels:
                    y = apply_label_nums(current_votes, labels, nums)
                    if y is not None:
                        overall_votes.append(y)
                elif nums:
                    overall_votes.append(nums[0])
                i = j
                continue

            if kind == 'overall' and label_num is not None:
                overall_votes.append(label_num)
            elif kind == 'overall' and i + 1 < len(lines) and standalone_num(lines[i + 1]) is not None:
                overall_votes.append(standalone_num(lines[i + 1]))
                i += 2
                continue
            elif kind in DIR_TITLES and label_num is not None:
                current_votes[kind].append(label_num)
            elif kind in DIR_TITLES and i + 1 < len(lines) and standalone_num(lines[i + 1]) is not None:
                current_votes[kind].append(standalone_num(lines[i + 1]))
                i += 2
                continue
            elif label_num is not None:
                for tid, title, needles in _CRITERIA:
                    if any(n in label_f for n in needles):
                        criteria.append({'id': tid, 'label': title, 'score': label_num})
                        break
            i += 1

    dirs = {did: pick_unique_score(nums) for did, nums in current_votes.items()}
    targets = {did: pick_unique_score(nums) for did, nums in target_votes.items()}
    yekun = pick_unique_score(overall_votes)
    if yekun is None:
        scored = [v for v in dirs.values() if v is not None]
        if len(scored) == 4:
            yekun = round(sum(scored) / 4, 1)
    return yekun, dirs, targets, criteria


def infer_dir_score(text, explicit):
    if explicit is not None:
        return explicit, False
    sents = sentences(text)
    if not sents:
        return None, True
    score = 48.0
    hits = 0
    for s in sents:
        pol = polarity(s)
        if pol == 'finding':
            score -= 7
            hits += 1
        elif pol == 'strength':
            score += 8
            hits += 1
        elif pol == 'action':
            score -= 2
            hits += 1
    if not hits:
        return 42.0, True
    return round(max(8.0, min(92.0, score)), 1), True


def theme_scores(text, criteria=None):
    seen = set()
    out = []
    for row in criteria or []:
        key = fold(row.get('label'))
        if not key or key in seen or row.get('score') is None:
            continue
        seen.add(key)
        out.append({
            'id': row.get('id') or key,
            'label': row.get('label'),
            'mentions': 1,
            'score': row.get('score'),
        })
    if out:
        out.sort(key=lambda x: -x['score'])
        return out[:8]
    f = fold(text)
    for tid, label, needles in THEMES:
        mentions = sum(f.count(n) for n in needles)
        if not mentions:
            continue
        chunk = ' '.join(s for s in sentences(text) if any(n in fold(s) for n in needles))
        pol = polarity(chunk or text)
        base = 40 + min(30, mentions * 6)
        if pol == 'strength':
            base += 18
        elif pol == 'finding':
            base -= 12
        out.append({
            'id': tid,
            'label': label,
            'mentions': mentions,
            'score': round(max(10, min(95, base)), 1),
        })
    out.sort(key=lambda x: (-x['mentions'], -x['score']))
    return out[:8]


def pick_unique(items, limit):
    seen = set()
    out = []
    for it in items:
        key = fold(it.get('text') if isinstance(it, dict) else it)
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(it)
        if len(out) >= limit:
            break
    return out


def classify_bits(slides):
    findings, strengths, actions = [], [], []
    for sl in slides:
        for sent in sentences((sl.get('title') or '') + '. ' + (sl.get('text') or '')):
            kind = polarity(sent)
            row = {'text': sent, 'dirId': dir_id_from_text(sent), 'slide': sl.get('n')}
            if kind == 'finding':
                findings.append(row)
            elif kind == 'strength':
                strengths.append(row)
            elif kind == 'action':
                actions.append(row)
    return (
        pick_unique(findings, 12),
        pick_unique(strengths, 10),
        pick_unique(actions, 10),
    )


def maturity_of(score):
    if score is None:
        return None
    if score < 25:
        return {'id': 'ilkin', 'label': 'İlkin'}
    if score < 50:
        return {'id': 'idare', 'label': 'İdarə olunan'}
    if score < 75:
        return {'id': 'mueyyen', 'label': 'Müəyyən edilmiş'}
    return {'id': 'opt', 'label': 'Optimallaşdırılan'}


def build_summary(org, overall, dirs, findings, strengths):
    name = org or 'Qurum'
    mat = maturity_of(overall)
    mat_lb = mat['label'] if mat else 'qiymətləndirmə'
    weak = None
    scored = [(did, dirs.get(did)) for did in DIR_TITLES if dirs.get(did) is not None]
    if scored:
        weak = min(scored, key=lambda x: x[1])
    parts = [
        '%s üzrə hesabat təhlili rəqəmsal inkişafı %s səviyyəsində (ümumi indeks %s/100) qiymətləndirir.'
        % (name, mat_lb.lower(), ('—' if overall is None else ('%s' % overall)))
    ]
    if weak:
        parts.append('Ən zəif istiqamət %s (%.1f bal) görünür.' % (DIR_TITLES[weak[0]], weak[1]))
    if findings:
        parts.append('Hesabatda %d əsas çatışmazlıq və boşluq qeyd olunub.' % len(findings))
    if strengths:
        parts.append('Eyni zamanda %d güclü tərəf təsbit edilib.' % len(strengths))
    return ' '.join(parts)


def name_from_file(filename):
    stem = os.path.splitext(os.path.basename(filename or ''))[0]
    cleaned = re.sub(
        r'(?i)diaqnostika|hesabat|reqemsallasma|pptx|pptm|pdf|prezentasiya|\bisq\b|\bexq\b',
        ' ',
        stem,
    )
    cleaned = re.sub(r'[_-]+', ' ', cleaned)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip(' ._-')
    return cleaned or stem or 'Hesabat'


def isq_id_from_label(label):
    f = fold(label)
    if not f:
        return ''
    if 'emeliyyat model' in f:
        return ''
    for pid, title, short, needles in ISQ_PILLARS:
        tf = fold(title)
        sf = fold(short)
        if f == tf or f == sf or (tf and tf in f) or any(n in f for n in needles):
            if pid == 'ops' and 'model' in f:
                continue
            if pid == 'infra' and 'texniki texnoloji' in f:
                continue
            return pid
    return ''


def detect_report_kind(filename, text):
    name = fold(filename or '')
    blob = fold((text or '')[:14000])
    hay = (name + ' ' + blob).strip()
    isq = exq = diag = 0
    if re.search(r'(^| )isq( |$)', name) or 'informasiya sistemlerinin qiymet' in name:
        isq += 14
    if re.search(r'(^| )exq( |$)', name) or 'elektron xidmetlerin qiymet' in name:
        exq += 14
    if 'diaqnostika' in name:
        diag += 14
    if 'diaqnostika' in blob:
        diag += 4
    if 'informasiya sistem' in blob:
        isq += 3
    if 'elektron xidmet' in blob and 'qiymet' in blob:
        exq += 4
    elif 'elektron xidmet' in blob:
        exq += 2
    if 'qerar 380' in hay or '4.14' in hay:
        exq += 3
        isq += 2
    if re.search(r'\b[1-5]\s*ulduz\b', blob) or 'ulduz sistemi' in blob:
        exq += 4
    isq_hits = 0
    for _pid, title, _short, needles in ISQ_PILLARS:
        if fold(title) in blob or any(n in blob for n in needles):
            isq_hits += 1
    if isq_hits >= 5:
        isq += 8
    elif isq_hits >= 3:
        isq += 4
    dir_hits = 0
    for did, needles in DIR_NEEDLES:
        if any(n in blob for n in needles):
            dir_hits += 1
    if dir_hits >= 3:
        diag += 5
    if 'reqemsal yetkinlik' in blob or 'reqemsallasma seviyyesi' in blob:
        diag += 6
    if 'movcud reqemsallasma' in blob:
        diag += 4
    ranked = sorted(
        (('diag', diag), ('isq', isq), ('exq', exq)),
        key=lambda x: x[1],
        reverse=True,
    )
    if ranked[0][1] <= 0:
        return 'diag'
    if ranked[0][1] == ranked[1][1]:
        if isq_hits >= 5:
            return 'isq'
        if dir_hits >= 3:
            return 'diag'
    return ranked[0][0]


def collect_overall_votes(lines):
    votes = []
    for i, line in enumerate(lines):
        folded = fold(line)
        label_f, label_num = split_label_score(line)
        if folded in _OVERALL_LABELS or 'yekun netice' in folded or 'umumi netice' in folded:
            if label_num is not None:
                votes.append(label_num)
            elif i + 1 < len(lines) and standalone_num(lines[i + 1]) is not None:
                votes.append(standalone_num(lines[i + 1]))
        m = re.search(r'(\d{1,3}(?:[.,]\d+)?)\s*/\s*100', line)
        if m and ('yekun' in folded or 'umumi' in folded or 'indeks' in folded or 'netice' in folded):
            n = parse_num(m.group(1))
            if n is not None:
                votes.append(n)
    return votes


def collect_isq_scores(pages):
    votes = {pid: [] for pid, _t, _s, _n in ISQ_PILLARS}
    overall_votes = []
    for sl in pages or []:
        lines = normalize_slide_lines((sl.get('title') or '') + '\n' + (sl.get('text') or ''))
        overall_votes.extend(collect_overall_votes(lines))
        i = 0
        while i < len(lines):
            line = lines[i]
            folded = fold(line)
            label_f, label_num = split_label_score(line)
            pid = isq_id_from_label(label_f or folded)
            if pid and label_num is not None:
                votes[pid].append(label_num)
            elif pid and i + 1 < len(lines) and standalone_num(lines[i + 1]) is not None:
                votes[pid].append(standalone_num(lines[i + 1]))
                i += 2
                continue
            i += 1
    dirs = {pid: pick_unique_score(nums) for pid, nums in votes.items()}
    yekun = pick_unique_score(overall_votes)
    scored = [v for v in dirs.values() if v is not None]
    if yekun is None and scored:
        yekun = round(sum(scored) / len(scored), 1)
    pillars = []
    for pid, title, short, _n in ISQ_PILLARS:
        pillars.append({
            'id': pid,
            'title': title,
            'short': short,
            'score': dirs.get(pid),
        })
    return yekun, dirs, pillars


def normalize_exq_mark(raw):
    if raw in (1, True, '1'):
        return 1
    if raw in (0, False, '0'):
        return 0
    s = str(raw or '').strip().upper().replace(' ', '')
    if s in ('NA', 'N/A', 'N\\A'):
        return 'NA'
    f = fold(raw)
    if 'shamil edilmir' in f or f in ('n a', 'na'):
        return 'NA'
    if 'temin edilir' in f:
        return 1
    if 'temin edilmir' in f or 'qismen' in f:
        return 0
    return None


def parse_exq_criteria(text):
    items = []
    seen = set()
    for m in re.finditer(
        r'(\d+\.\d+\.\d+)(?:\s*[|\t;:,\-]\s*|\s+)(?:([1-5])(?:\s*[|\t;:,\-]\s*|\s+))?(1|0|N/A|NA|n/a)\b',
        text or '',
        re.I,
    ):
        mark = normalize_exq_mark(m.group(3))
        star = int(m.group(2) or 0) or None
        if star is None:
            around = (text or '')[max(0, m.start()):min(len(text or ''), m.end() + 28)]
            sm = re.search(r'([1-5])\s*ulduz', fold(around))
            if sm:
                star = int(sm.group(1))
        key = (m.group(1), star, mark)
        if key in seen or mark is None:
            continue
        seen.add(key)
        items.append({'code': m.group(1), 'star': star, 'value': mark})
    for m in re.finditer(
        r'(\d+\.\d+\.\d+)[^\n]{0,80}?\b([1-5])\s*ulduz\b[^\n]{0,40}?\b(1|0|N/A|NA)\b',
        text or '',
        re.I,
    ):
        mark = normalize_exq_mark(m.group(3))
        key = (m.group(1), int(m.group(2)), mark)
        if key in seen or mark is None:
            continue
        seen.add(key)
        items.append({'code': m.group(1), 'star': int(m.group(2)), 'value': mark})
    return items


def compute_exq_yekun(items):
    collected = [0, 0, 0, 0, 0]
    possible = [0, 0, 0, 0, 0]
    applicable = 0
    for it in items or []:
        try:
            star = int(it.get('star'))
        except (TypeError, ValueError):
            continue
        if star < 1 or star > 5:
            continue
        val = it.get('value')
        if val in ('NA', 'N/A'):
            continue
        met = val in (1, '1', True)
        missed = val in (0, '0', False)
        if not met and not missed:
            continue
        applicable += 1
        for k in range(1, 6):
            if star < k:
                continue
            possible[k - 1] += 1
            if met:
                collected[k - 1] += 1
    tot_c = tot_p = 0.0
    by_star = []
    for i in range(5):
        tot_c += collected[i] * EXQ_STAR_WEIGHTS[i]
        tot_p += possible[i] * EXQ_STAR_WEIGHTS[i]
        by_star.append({
            'star': i + 1,
            'weight': EXQ_STAR_WEIGHTS[i],
            'collected': collected[i],
            'possible': possible[i],
        })
    if tot_p <= 0:
        return {'percent': None, 'star': None, 'applicable': applicable, 'byStar': by_star}
    percent = round((tot_c / tot_p) * 100, 1)
    if percent < 21:
        star = 1
    elif percent < 41:
        star = 2
    elif percent < 61:
        star = 3
    elif percent < 81:
        star = 4
    else:
        star = 5
    return {'percent': percent, 'star': star, 'applicable': applicable, 'byStar': by_star}


def collect_exq_scores(pages):
    full = '\n'.join((p.get('title') or '') + '\n' + (p.get('text') or '') for p in (pages or []))
    items = parse_exq_criteria(full)
    computed = compute_exq_yekun(items)
    overall_votes = []
    star_votes = []
    for sl in pages or []:
        lines = normalize_slide_lines((sl.get('title') or '') + '\n' + (sl.get('text') or ''))
        overall_votes.extend(collect_overall_votes(lines))
        for line in lines:
            f = fold(line)
            sm = re.search(r'\b([1-5])\s*ulduz\b', f)
            if sm and ('yekun' in f or 'netice' in f or 'seviyye' in f):
                star_votes.append(int(sm.group(1)))
            m = re.search(r'(\d{1,3}(?:[.,]\d+)?)\s*%', line)
            if m and ('yekun' in f or 'umumi' in f):
                n = parse_num(m.group(1))
                if n is not None:
                    overall_votes.append(n)
    yekun = pick_unique_score(overall_votes)
    if yekun is None:
        yekun = computed.get('percent')
    star = pick_unique_score(star_votes)
    if star is None:
        star = computed.get('star')
    elif yekun is None:
        bands = ((1, 10), (2, 30), (3, 50), (4, 70), (5, 90))
        yekun = dict(bands).get(int(star))
    return yekun, star, items, computed.get('byStar') or []


def dir_id_for_kind(text, kind):
    if kind == 'isq':
        return isq_id_from_label(text) or dir_id_from_text(text)
    return dir_id_from_text(text)


def classify_bits_kind(slides, kind):
    findings, strengths, actions = [], [], []
    for sl in slides:
        for sent in sentences((sl.get('title') or '') + '. ' + (sl.get('text') or '')):
            pol = polarity(sent)
            row = {
                'text': sent,
                'dirId': dir_id_for_kind(sent, kind),
                'slide': sl.get('n'),
            }
            if pol == 'finding':
                findings.append(row)
            elif pol == 'strength':
                strengths.append(row)
            elif pol == 'action':
                actions.append(row)
    return (
        pick_unique(findings, 12),
        pick_unique(strengths, 10),
        pick_unique(actions, 10),
    )


def build_kind_summary(kind, org, overall, dirs, findings, strengths, star=None):
    name = org or 'Qurum'
    if kind == 'isq':
        weak = None
        scored = [(pid, dirs.get(pid), title) for pid, title, _s, _n in ISQ_PILLARS if dirs.get(pid) is not None]
        if scored:
            weak = min(scored, key=lambda x: x[1])
        parts = [
            '%s üzrə İSQ hesabatı yekun nəticəni %s qiymətləndirir.'
            % (name, ('—' if overall is None else ('%s/100' % overall)))
        ]
        if weak:
            parts.append('Ən zəif sütun %s (%.1f).' % (weak[2], weak[1]))
    elif kind == 'exq':
        star_lb = ('%s ulduz' % star) if star else 'ulduz şkalası'
        parts = [
            '%s üzrə EXQ hesabatı yekun nəticəni %s (%s) göstərir.'
            % (name, ('—' if overall is None else ('%s%%' % overall)), star_lb)
        ]
    else:
        return build_summary(org, overall, dirs, findings, strengths)
    if findings:
        parts.append('Hesabatda %d əsas çatışmazlıq qeyd olunub.' % len(findings))
    if strengths:
        parts.append('%d güclü tərəf təsbit edilib.' % len(strengths))
    return ' '.join(parts)


def empty_kind_fields():
    return {
        'pillars': [],
        'star': None,
        'byStar': [],
        'criteria': [],
    }


def normalize_report_kind(kind):
    k = fold(kind)
    if k in ('isq',):
        return 'isq'
    if k in ('exq',):
        return 'exq'
    if k in ('diag', 'diaqnostika', 'diagnostika'):
        return 'diag'
    return ''


def analyze_slides(slides, filename='', kind=None):
    blobs = [(s.get('title') or '') + '\n' + (s.get('text') or '') for s in slides]
    full = '\n'.join(blobs)
    org = org_from_text(blobs) or name_from_file(filename)
    year = year_from_text(full)
    kind = normalize_report_kind(kind) or detect_report_kind(filename, full)
    extra = empty_kind_fields()
    inferred_any = False
    targets = {did: None for did in DIR_TITLES}
    crit_rows = []
    star = None

    if kind == 'isq':
        yekun, dirs, pillars = collect_isq_scores(slides)
        extra['pillars'] = pillars
        findings, strengths, actions = classify_bits_kind(slides, 'isq')
        summary = build_kind_summary('isq', org, yekun, dirs, findings, strengths)
        themes = theme_scores(full)
        if yekun is None:
            inferred_any = False
    elif kind == 'exq':
        yekun, star, items, by_star = collect_exq_scores(slides)
        extra['star'] = star
        extra['byStar'] = by_star
        extra['criteria'] = items
        dirs = {}
        findings, strengths, actions = classify_bits_kind(slides, 'exq')
        summary = build_kind_summary('exq', org, yekun, dirs, findings, strengths, star=star)
        themes = theme_scores(full)
    else:
        kind = 'diag'
        yekun, explicit_dirs, targets, crit_rows = collect_scores(slides)
        dirs = {}
        for did in DIR_TITLES:
            related = '\n'.join(
                b for b in blobs
                if dir_id_from_text(b) == did or any(n in fold(b) for n in dict(DIR_NEEDLES)[did])
            )
            score, inferred = infer_dir_score(related or full, explicit_dirs.get(did))
            dirs[did] = score
            inferred_any = inferred_any or inferred
        scored = [v for v in dirs.values() if v is not None]
        if yekun is None and scored:
            yekun = round(sum(scored) / len(scored), 1)
            inferred_any = True
        findings, strengths, actions = classify_bits(slides)
        themes = theme_scores(full, crit_rows)
        summary = build_summary(org, yekun, dirs, findings, strengths)
        inferred_any = inferred_any and yekun is not None and all(
            explicit_dirs.get(d) is None for d in DIR_TITLES
        )

    outline = []
    for s in slides:
        title = clean_slide_title(s.get('title') or '')
        outline.append({
            'n': s.get('n'),
            'title': title or ('Slayd %s' % s.get('n')),
        })
    if not actions and findings:
        for row in findings[:4]:
            did = row.get('dirId') or ''
            if kind == 'isq':
                title = next((t for pid, t, _s, _n in ISQ_PILLARS if pid == did), 'İSQ')
            elif kind == 'exq':
                title = 'elektron xidmət'
            else:
                title = DIR_TITLES.get(did, 'rəqəmsallaşma')
            actions.append({
                'text': '%s üzrə növbəti rəsmi səviyyəyə çatmaq üçün tədbirlər planı hazırlanıb icra edilsin.' % title,
                'dirId': did,
                'slide': row.get('slide'),
            })
    out = {
        'kind': kind,
        'org': org,
        'year': year,
        'slideCount': len(slides),
        'pageCount': len(slides),
        'overall': yekun,
        'inferred': bool(inferred_any),
        'maturity': maturity_of(yekun),
        'dirs': dirs,
        'targets': {did: (targets or {}).get(did) for did in DIR_TITLES},
        'summary': summary,
        'findings': findings,
        'strengths': strengths,
        'actions': actions,
        'themes': themes,
        'slides': outline,
    }
    out.update(extra)
    return out


def parse_report_file(path, filename=None, kind=None):
    filename = filename or os.path.basename(path)
    ext = os.path.splitext(filename)[1].lower()
    if ext in ('.pptx', '.pptm'):
        pages = extract_pptx_slides(path)
        fmt = 'pptx'
        empty_msg = 'Təqdimatda oxunaqlı mətn tapılmadı.'
    elif ext == '.pdf':
        pages = extract_pdf_pages(path)
        fmt = 'pdf'
        empty_msg = 'PDF-də oxunaqlı mətn tapılmadı. Skan olunmuş səhifələr tanınmır.'
    else:
        raise RuntimeError('Yalnız .pptx, .pptm və .pdf qəbul olunur.')
    if not pages or not any((s.get('text') or '').strip() for s in pages):
        return {'report': None, 'warnings': [empty_msg]}
    report = analyze_slides(pages, filename, kind=kind)
    report['format'] = fmt
    return {'report': report, 'warnings': []}


def parse_report_pptx(path, filename=None):
    return parse_report_file(path, filename)
