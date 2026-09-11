# -*- coding: utf-8 -*-
"""Qurum diaqnostikası Excel-i: istiqamət→pillar, meyar→domain, alt-meyar→capability."""
from __future__ import annotations

import csv
import io
import os
import re
from collections import OrderedDict
from datetime import date, datetime


def require_openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError(
            'Excel oxumaq üçün openpyxl lazımdır. Terminalda: pip install -r requirements.txt'
        ) from None
    return load_workbook

DIR_NEEDLES = (
    ('strategiya', ('strategiya', 'strategy', 'strategic')),
    ('xidmetler', ('xidmet', 'service')),
    ('texniki', ('texniki', 'infrastruktur', 'texnoloj', 'technical', 'technolog')),
    ('emeliyyat', ('emeliyyat', 'meliyyat', 'operation', 'operating')),
)

DIR_TITLES = {
    'strategiya': 'Strategiya',
    'xidmetler': 'Xidmətlər',
    'texniki': 'Texniki-texnoloji infrastruktur',
    'emeliyyat': 'Əməliyyat modelləri',
}

HEADER_MAP = {
    'org': ('qurum', 'teskilat', 'muessise', 'institution', 'teskilatin adi', 'qurumun adi'),
    'dir': ('istiqamet', 'direction', 'pillar', 'pillars'),
    'crit': ('meyar', 'criterion', 'criteria', 'domain', 'domains'),
    'sub': ('altmeyar', 'alt meyar', 'subcriterion', 'sub criterion', 'capability', 'capabilities'),
    'current': (
        'cari veziyyet', 'cari veziyyət', 'hazirki veziyyet', 'hazirki veziyyət',
        'movcud veziyyet', 'current', 'veziyyet',
        'aciqlama', 'achiqlama', 'tesvir', 'izah', 'description', 'sherh',
        'netice aciqlama',
    ),
    'gap': (
        'catismaz', 'bosluq', 'deficiency', 'gap', 'uygunsuz',
        'tapilan problem', 'musahide'
    ),
    'score': ('bal', 'qiymet', 'qiymetlendirme', 'netice', 'result', 'score', 'qrsg', 'yekun'),
    'year': ('il', 'year', 'dovr', 'tarix'),
}

KV_LABELS = {
    'org': ('qurum', 'teskilat', 'muessise', 'diaqnostika aparilan'),
    'year': ('il', 'dovr', 'tarix'),
    'score': ('qrsg', 'umumi netice', 'yekun netice', 'reqemsallasma seviyyesi', 'umumi bal'),
    'current': ('cari veziyyet', 'hazirki veziyyet'),
    'gap': ('catismaz', 'bosluq'),
}

SKIP_HEADER = (
    'movcud veziyyet ve qiymet',
    'ballarla hesablanir',
    'sira no',
    'sira n',
)

OFFICIAL_SCORES = (0, 50, 75, 100)
LEVEL_TO_SCORE = {1: 0.0, 2: 33.3, 3: 66.6, 4: 100.0}
LEVEL_WORDS = {
    1: ('ilkin', 'initial'),
    2: ('idare', 'managed'),
    3: ('mueyyen', 'defined'),
    4: ('optim',),
}
STRUCT_HEADER = frozenset({'dir', 'crit', 'sub'})
USEFUL_HEADER = STRUCT_HEADER | frozenset({'current', 'gap', 'score', 'org'})
MARK_TOKENS = frozenset({'x', 'v', '1', 'yes', 'beli', 'ok', '+'})
NO_GAP = 'Bu altmeyar üzrə əsas çatışmazlıq aşkar edilməyib.'
DIR_GAP_HINTS = {
    'strategiya': 'Rəqəmsallaşma üzrə strateji baxış və icra mexanizmi formalaşdırılmayıb və ya ilkin səviyyədədir.',
    'xidmetler': 'Xidmət prosesləri əsasən kağız və əl üsulu ilə aparılır; rəqəmsal izləmə zəifdir.',
    'texniki': 'Texniki-texnoloji infrastruktur inventarlaşdırılmayıb və ya standartlar müəyyən edilməyib.',
    'emeliyyat': 'Əməliyyat modelləri sənədləşdirilməyib; proseslər rəqəmsallaşdırılmayıb.',
}

_TR = str.maketrans({
    'ı': 'i', 'İ': 'i', 'I': 'i', 'ə': 'e', 'Ə': 'e',
    'ö': 'o', 'Ö': 'o', 'ü': 'u', 'Ü': 'u',
    'ğ': 'g', 'Ğ': 'g', 'ş': 's', 'Ş': 's', 'ç': 'c', 'Ç': 'c',
})


def fold(value):
    s = '' if value is None else str(value)
    s = s.replace('\xa0', ' ').translate(_TR).lower()
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return ''
    if isinstance(value, datetime):
        return value.strftime('%d.%m.%Y')
    if isinstance(value, date):
        return value.strftime('%d.%m.%Y')
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value != int(value):
            return str(value)
        return str(int(value)) if float(value).is_integer() else str(value)
    s = str(value).replace('\xa0', ' ').strip()
    if s.lower() in ('none', 'null', '-', '—', 'n/a'):
        return ''
    return s


def parse_num(value):
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        n = float(value)
        return n if n == n and abs(n) <= 1000 else None
    s = str(value).strip().replace(',', '.')
    m = re.search(r'-?\d+(?:\.\d+)?', s)
    if not m:
        return None
    n = float(m.group(0))
    return n if abs(n) <= 1000 else None


def dir_id_from_title(title):
    f = fold(title)
    if not f:
        return ''
    if 'xidmet' in f and 'say' in f:
        return ''
    for did, needles in DIR_NEEDLES:
        for n in needles:
            if n in f:
                return did
    return ''


def header_hit(folded, needle):
    if not folded or not needle:
        return False
    if len(needle) <= 3:
        return folded == needle or bool(re.search(r'(^|\s)' + re.escape(needle) + r'($|\s)', folded))
    return needle in folded


def header_role(text):
    f = fold(text)
    if not f or any(tok in f for tok in SKIP_HEADER):
        return ''
    if f in ('0', '50', '75', '100') or re.fullmatch(r'\d+', f):
        return ''
    if len(f) > 48:
        return ''
    best = ''
    best_len = 0
    for role, needles in HEADER_MAP.items():
        for n in needles:
            if header_hit(f, n) and len(n) > best_len:
                best = role
                best_len = len(n)
    if best == 'current' and 'qiymet' in f and 'veziyyet' in f:
        return ''
    if best == 'score' and ('cari' in f or 'catismaz' in f) and 'bal' not in f and 'qiymet' not in f:
        return ''
    if best in ('sub', 'crit', 'dir') and any(header_hit(f, n) for n in HEADER_MAP['score']):
        return 'score'
    return best


def official_header_mark(text):
    f = fold(text)
    if not f:
        return None
    if f in ('0', '50', '75', '100'):
        return int(f)
    m = re.match(r'^(0|50|75|100)(?:\s+|$)', f)
    if m and len(f) <= 40:
        return int(m.group(1))
    return None


def level_header_mark(text):
    f = fold(text)
    if not f:
        return None
    m = re.match(r'^([1-4])(?:\s+|$)', f)
    if not m:
        return None
    n = int(m.group(1))
    if any(tok in f for tok in LEVEL_WORDS[n]):
        return n
    return None


def score_header_priority(text):
    f = fold(text)
    if 'netice' in f or f == 'result':
        return 10
    if 'yekun' in f or 'qrsg' in f:
        return 8
    if 'iddia' in f:
        return 1
    if 'bal' in f and 'qiymet' not in f:
        return 3
    return 5


def is_mark_token(text):
    raw = cell_text(text)
    f = fold(raw)
    if not f:
        return False
    if f in MARK_TOKENS or f in ('0', '50', '75', '100'):
        return True
    return '✓' in raw or '✔' in raw


def is_no_gap(text):
    f = fold(text)
    return 'askar edilmeyib' in f or 'ashkar edilmeyib' in f


def split_score_text(value):
    text = cell_text(value)
    n = parse_num(value)
    if not text:
        return n, ''
    m = re.match(
        r'^\s*(-?\d+(?:[.,]\d+)?)\s*(?:/\s*\d+(?:[.,]\d+)?)?\s*[-–—:.)]?\s*(.*)$',
        text,
    )
    if m:
        rest = (m.group(2) or '').strip()
        if rest and fold(rest) not in ('bal', 'xal', '%', 'ball'):
            return n, rest
        return n, ''
    if n is not None:
        return n, ''
    return None, text


def nearest_official(score):
    if score is None:
        return None
    return min(OFFICIAL_SCORES, key=lambda s: (abs(s - score), -s))


def next_official_band(score):
    if score is None:
        return None
    for s in OFFICIAL_SCORES:
        if s > score:
            return s
    return None


def rubric_texts(row, marks, levels=None):
    out = {}
    pairs = list((marks or {}).items()) + list((levels or {}).items())
    for key, idx in pairs:
        if idx >= len(row):
            continue
        text = cell_text(row[idx])
        if not text or is_mark_token(text):
            continue
        compact = fold(text).replace(' ', '')
        if compact in ('0', '50', '75', '100') or re.fullmatch(r'-?\d+(?:[.,]\d+)?', compact):
            continue
        out[int(key)] = text.strip()
    return out


def used_indexes(cols, marks):
    used = set()
    for v in (cols or {}).values():
        used.add(v)
    for v in (marks or {}).values():
        used.add(v)
    return used


def neighbor_explain(row, cols, marks):
    score_idx = (cols or {}).get('score')
    if score_idx is None:
        return ''
    used = used_indexes(cols, marks)
    for idx in (score_idx + 1, score_idx - 1):
        if idx < 0 or idx >= len(row) or idx in used:
            continue
        text = cell_text(row[idx])
        if not text or len(text) < 8:
            continue
        if official_header_mark(text) is not None:
            continue
        if header_role(text):
            continue
        if parse_year(text):
            continue
        return text
    return ''


def pick_current(score, current, rubrics, level=None):
    current = (current or '').strip()
    if current:
        return current
    if not rubrics:
        return ''
    if level is not None and int(level) in rubrics:
        return rubrics[int(level)]
    if score is None:
        return ''
    exact = int(score) if float(score) == int(score) else None
    if exact in rubrics:
        return rubrics[exact]
    band = nearest_official(score)
    if band in rubrics and abs(band - score) <= 25:
        return rubrics[band]
    return ''


def band_hint(score, dir_id):
    if score is None:
        return ''
    if score < 25:
        return DIR_GAP_HINTS.get(
            dir_id,
            'Bu altmeyar ilkin səviyyədə qiymətləndirilib; əsas proseslər qurulmayıb.',
        )
    if score < 50:
        return 'İdarə olunan səviyyədə qalıb; vahid yanaşma və nəzarət mexanizmi tam formalaşmayıb.'
    if score < 75:
        return 'Müəyyən edilmiş səviyyəyə çatıb, lakin ölçmə, hesabatlılıq və davamlı təkmilləşdirmə tam təmin olunmayıb.'
    if score < 100:
        return 'Optimallaşdırılan səviyyəyə tam keçid edilməyib; nəticələr mütəmadi ölçülərək təkmilləşdirilməlidir.'
    return NO_GAP


def infer_gap(score, current, rubrics, title, dir_id, level=None):
    rubrics = rubrics or {}
    if level is not None:
        if level >= 4 or (score is not None and score >= 100):
            return NO_GAP
        nxt = int(level) + 1
        if rubrics.get(nxt):
            return 'Növbəti səviyyəyə (%s) çatmayıb: %s' % (nxt, rubrics[nxt])
        return band_hint(score if score is not None else LEVEL_TO_SCORE.get(level), dir_id)
    if score is None:
        return ''
    if score >= 100:
        return NO_GAP
    nxt = next_official_band(score)
    if nxt is not None and rubrics.get(nxt):
        return 'Növbəti səviyyəyə (%s) çatmayıb: %s' % (nxt, rubrics[nxt])
    return band_hint(score, dir_id)


def claimed_score(row, cols, marks, levels):
    used = used_indexes(cols, marks)
    for idx in (levels or {}).values():
        used.add(idx)
    found = []
    for i, cell in enumerate(row or []):
        if i in used:
            continue
        text = cell_text(cell)
        if not text:
            continue
        n = parse_num(text)
        if n is None or n > 100 or n < 0:
            continue
        compact = fold(text).replace(' ', '')
        if compact in ('1', '2', '3', '4'):
            continue
        found.append(n)
    return found[0] if len(found) == 1 else None


def resolve_row_score(row, cols, marks, levels=None):
    levels = levels or {}
    score = marked_score(row, marks)
    raw = take(row, cols, 'score')
    sn, stext = split_score_text(raw)
    level = None
    if levels and sn is not None and sn in (1, 2, 3, 4) and float(sn) == int(sn):
        level = int(sn)
        claimed = claimed_score(row, cols, marks, levels)
        if claimed is not None:
            score = claimed
        elif score is None:
            score = LEVEL_TO_SCORE[level]
        stext = ''
    elif score is None:
        score = sn
    if score is not None and score > 100:
        score = None
    return score, stext, level


def collect_current(row, cols, marks, score, score_text, levels=None, level=None):
    current = take(row, cols, 'current')
    if current:
        return current
    if score_text:
        return score_text
    picked = pick_current(score, '', rubric_texts(row, marks, levels), level)
    if picked:
        return picked
    return neighbor_explain(row, cols, marks)


def merged_fill_map(ws):
    filled = {}
    for rng in ws.merged_cells.ranges:
        val = ws.cell(rng.min_row, rng.min_col).value
        if val is None:
            continue
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                if row == rng.min_row and col == rng.min_col:
                    continue
                filled[(row, col)] = val
    return filled


def sheet_matrix(ws):
    fills = merged_fill_map(ws)
    rows = []
    for r_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = []
        for c_i, cell in enumerate(row, start=1):
            if cell is None and (r_i, c_i) in fills:
                cell = fills[(r_i, c_i)]
            vals.append(cell_text(cell))
        while vals and not vals[-1]:
            vals.pop()
        rows.append(vals)
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def look_kv(rows, max_rows=30):
    meta = {}
    for row in rows[:max_rows]:
        if len(row) < 2:
            continue
        left = fold(row[0])
        right = row[1]
        if not left or not right:
            continue
        for role, needles in KV_LABELS.items():
            if any(n in left for n in needles) and role not in meta:
                meta[role] = right
                break
    return meta


def row_header_maps(row):
    colmap = {}
    marks = {}
    levels = {}
    score_hits = []
    for c, cell in enumerate(row or []):
        mk = official_header_mark(cell)
        if mk is not None:
            marks[mk] = c
            continue
        lv = level_header_mark(cell)
        if lv is not None:
            levels[lv] = c
            continue
        role = header_role(cell)
        if role == 'score':
            score_hits.append((score_header_priority(cell), c))
            continue
        if role and role not in colmap:
            colmap[role] = c
    if score_hits:
        score_hits.sort()
        colmap['score'] = score_hits[-1][1]
    if len(marks) == 1 and not levels:
        marks = {}
    return colmap, marks, levels


def looks_like_subheader(row):
    if not any(row):
        return False
    texts = [cell_text(c) for c in row]
    if any(len(t) > 40 for t in texts):
        return False
    cols, marks, levels = row_header_maps(texts)
    if not marks and not levels and not cols:
        return False
    for t in texts:
        if not t:
            continue
        if official_header_mark(t) is not None or level_header_mark(t) is not None:
            continue
        if header_role(t):
            continue
        return False
    return True


def header_quality(colmap, marks, levels=None):
    useful_n = len(USEFUL_HEADER.intersection(colmap))
    struct_n = len(STRUCT_HEADER.intersection(colmap))
    bands = bool(marks) or bool(levels)
    if useful_n < 2 and not (bands and struct_n >= 2):
        return 0
    return useful_n + (2 if bands else 0) + struct_n + (2 if levels else 0)


def find_header(rows):
    best_i = -1
    best_map = {}
    best_score = 0
    limit = min(40, len(rows))
    for i in range(limit):
        cols, marks, levels = row_header_maps(rows[i])
        header_end = i
        if i + 1 < len(rows) and looks_like_subheader(rows[i + 1]):
            cols2, marks2, levels2 = row_header_maps(rows[i + 1])
            merged_cols = dict(cols)
            merged_marks = dict(marks)
            merged_levels = dict(levels)
            for role, idx in cols2.items():
                if role == 'score':
                    continue
                merged_cols.setdefault(role, idx)
            if 'score' in cols2:
                prev = cols.get('score')
                if prev is None or score_header_priority(rows[i + 1][cols2['score']]) > score_header_priority(rows[i][prev] if prev < len(rows[i]) else ''):
                    merged_cols['score'] = cols2['score']
            for mk, idx in marks2.items():
                merged_marks.setdefault(mk, idx)
            for lv, idx in levels2.items():
                merged_levels.setdefault(lv, idx)
            added = (
                len(merged_cols) > len(cols)
                or (merged_marks and not marks)
                or len(merged_marks) > len(marks)
                or (merged_levels and not levels)
            )
            if added:
                cols, marks, levels = merged_cols, merged_marks, merged_levels
                header_end = i + 1
        q = header_quality(cols, marks, levels)
        if q > best_score:
            best_score = q
            best_i = header_end
            best_map = {'cols': cols, 'marks': marks, 'levels': levels}
    if best_i < 0:
        return -1, {}, {}, {}
    return best_i, best_map['cols'], best_map['marks'], best_map['levels']


def marked_score(row, marks):
    found = []
    for score, idx in (marks or {}).items():
        if idx >= len(row):
            continue
        raw = row[idx]
        f = fold(raw)
        if not f:
            continue
        if f in MARK_TOKENS or '✓' in (raw or '') or '✔' in (raw or ''):
            found.append(score)
            continue
        if f.replace(' ', '') == str(int(score)):
            found.append(score)
    if len(found) == 1:
        return float(found[0])
    return None


def take(row, cols, role):
    idx = cols.get(role)
    if idx is None or idx >= len(row):
        return ''
    return row[idx].strip()


def empty_org(name, year=None):
    return {
        'name': name or '',
        'year': year,
        'date': None,
        'qrsg': None,
        'dirs': {d: None for d in DIR_TITLES},
        'dirTexts': {d: '' for d in DIR_TITLES},
        'dirGaps': {d: '' for d in DIR_TITLES},
        'extras': [],
        'overallText': '',
        'overallCurrent': '',
        'overallGaps': '',
    }


def ensure_org(bucket, name, year):
    key = name.strip() or '__default__'
    if key not in bucket:
        bucket[key] = empty_org(name.strip(), year)
    org = bucket[key]
    if year and org['year'] is None:
        org['year'] = year
    return org


def append_unique(parts, text):
    t = (text or '').strip()
    if not t:
        return
    if t not in parts:
        parts.append(t)


def add_extra(org, dir_id, title, score, current, gap, criterion='', sub=''):
    title = (title or '').strip()
    current = (current or '').strip()
    gap = (gap or '').strip()
    criterion = (criterion or '').strip()
    sub = (sub or '').strip()
    if not title and not criterion and not sub and score is None and not current and not gap:
        return
    if not title:
        title = sub or criterion or DIR_TITLES.get(dir_id) or 'Meyar'
    blob = current or gap
    for prev in org['extras']:
        same = (
            (prev.get('dirId') or '') == (dir_id or '')
            and fold(prev.get('criterion') or '') == fold(criterion)
            and fold(prev.get('sub') or prev.get('title') or '') == fold(sub or title)
        )
        if same:
            if score is not None and prev.get('score') is None:
                prev['score'] = score
            if current and not prev.get('current'):
                prev['current'] = current
            if gap and not prev.get('deficiency'):
                prev['deficiency'] = gap
            if blob and not prev.get('text'):
                prev['text'] = blob
            if criterion and not prev.get('criterion'):
                prev['criterion'] = criterion
            if sub and not prev.get('sub'):
                prev['sub'] = sub
            return
    org['extras'].append({
        'title': title,
        'criterion': criterion,
        'sub': sub,
        'score': score,
        'text': blob,
        'current': current,
        'deficiency': gap,
        'dirId': dir_id or '',
    })


def apply_dir_level(org, dir_id, score, current, gap):
    if not dir_id:
        return
    if score is not None and org['dirs'].get(dir_id) is None:
        org['dirs'][dir_id] = score
    if current:
        prev = org['dirTexts'].get(dir_id) or ''
        org['dirTexts'][dir_id] = (prev + '\n' + current).strip() if prev and current not in prev else (prev or current)
    if gap:
        prev = org['dirGaps'].get(dir_id) or ''
        org['dirGaps'][dir_id] = (prev + '\n' + gap).strip() if prev and gap not in prev else (prev or gap)


def parse_table(rows, header_i, cols, marks, sheet_dir, meta, bucket, levels=None):
    levels = levels or {}
    default_name = cell_text(meta.get('org'))
    default_year = parse_year(meta.get('year'))
    last_org = default_name
    last_year = default_year
    last_dir_raw = ''
    last_crit = ''
    last_did = sheet_dir
    for row in rows[header_i + 1:]:
        if not any(row):
            continue
        joined = fold(' '.join(cell_text(c) for c in row[:8]))
        header_like = sum(1 for c in row if header_role(c) and len(cell_text(c)) <= 22)
        if header_like >= 3 and not any(len(cell_text(c)) > 40 for c in row):
            continue
        if 'istiqamet' in joined and 'meyar' in joined:
            continue
        if 'pillar' in joined and 'domain' in joined:
            continue
        explicit_org = take(row, cols, 'org')
        if explicit_org:
            last_org = explicit_org
            last_dir_raw = ''
            last_crit = ''
            last_did = sheet_dir
        org_name = explicit_org or last_org or default_name
        year_raw = take(row, cols, 'year')
        year = parse_year(year_raw) or last_year or default_year
        if parse_year(year_raw):
            last_year = year
        org = ensure_org(bucket, org_name, year)
        explicit_dir = take(row, cols, 'dir')
        explicit_crit = take(row, cols, 'crit')
        sub = take(row, cols, 'sub')
        score, score_text, level = resolve_row_score(row, cols, marks, levels)
        rubrics = rubric_texts(row, marks, levels)
        current = collect_current(row, cols, marks, score, score_text, levels, level)
        gap = take(row, cols, 'gap')
        if explicit_dir and not explicit_crit and not sub and score is None and not current and not gap:
            last_dir_raw = explicit_dir
            last_did = dir_id_from_title(explicit_dir) or last_did
            last_crit = ''
            continue
        if explicit_crit and not sub and score is None and not current and not gap:
            if explicit_dir:
                last_dir_raw = explicit_dir
                last_did = dir_id_from_title(explicit_dir) or last_did
            last_crit = explicit_crit
            continue
        if explicit_dir:
            last_dir_raw = explicit_dir
            if not explicit_crit:
                last_crit = ''
        if explicit_crit:
            last_crit = explicit_crit
        dir_raw = explicit_dir or last_dir_raw
        crit = explicit_crit or last_crit
        did = dir_id_from_title(dir_raw) or sheet_dir
        if not did and not dir_raw:
            did = dir_id_from_title(crit)
        if did:
            last_did = did
        else:
            did = last_did
        title = sub or crit or dir_raw
        if score is None and not current and not gap:
            continue
        if not gap:
            gap = infer_gap(score, current, rubrics, title, did, level)
        if not did and not title and score is None and not current and not gap:
            continue
        if (sub or crit) and did:
            add_extra(org, did, title, score, current, gap, crit, sub)
        elif did and not sub and not crit:
            apply_dir_level(org, did, score, current, gap)
        else:
            add_extra(org, did, title or 'Meyar', score, current, gap, crit, sub)


def parse_pairs(rows, sheet_dir, meta, bucket):
    org = ensure_org(bucket, cell_text(meta.get('org')), parse_year(meta.get('year')))
    did = sheet_dir
    title = ''
    criterion = ''
    sub = ''
    current_parts = []
    gap_parts = []
    score = None

    def flush():
        nonlocal title, current_parts, gap_parts, score, sub
        cur = '\n'.join(current_parts).strip()
        gp = '\n'.join(gap_parts).strip()
        if not gp and score is not None:
            gp = infer_gap(score, cur, {}, title or sub or criterion, did)
        if title or criterion or sub or cur or gp or score is not None:
            if title or criterion or sub:
                add_extra(org, did, title or sub or criterion, score, cur, gp, criterion, sub)
            elif did:
                apply_dir_level(org, did, score, cur, gp)
            else:
                if cur:
                    org['overallCurrent'] = (org['overallCurrent'] + '\n' + cur).strip() if org['overallCurrent'] else cur
                if gp:
                    org['overallGaps'] = (org['overallGaps'] + '\n' + gp).strip() if org['overallGaps'] else gp
        title = ''
        sub = ''
        current_parts = []
        gap_parts = []
        score = None

    for row in rows:
        if not any(row):
            continue
        left = fold(row[0] if row else '')
        right = row[1] if len(row) > 1 else ''
        if not left:
            continue
        nd = dir_id_from_title(row[0])
        if nd and (left.startswith('istiqamet') or nd and len(left) < 40):
            if 'istiqamet' in left or dir_id_from_title(row[0]):
                if not right:
                    flush()
                    did = nd or did
                    continue
        role = header_role(row[0])
        if role == 'dir' or 'istiqamet' in left or left.startswith('pillar'):
            flush()
            did = dir_id_from_title(right) or nd or did
            continue
        if role == 'crit' or 'domain' in left or left.startswith('meyar'):
            flush()
            criterion = right or cell_text(row[0])
            title = criterion
            continue
        if role == 'sub' or 'capabilit' in left or left.startswith('altmeyar'):
            flush()
            sub = right or cell_text(row[0])
            title = sub
            continue
        if role == 'current' or left.startswith('cari') or 'veziyyet' in left:
            append_unique(current_parts, right or ' '.join(row[1:]))
            continue
        if role == 'gap' or 'catismaz' in left or left.startswith('bosluq'):
            append_unique(gap_parts, right or ' '.join(row[1:]))
            continue
        if role == 'score' or 'bal' in left or 'qiymet' in left or 'netice' in left:
            sn, stext = split_score_text(right)
            if sn is not None:
                score = sn
            if stext:
                append_unique(current_parts, stext)
            continue
        if nd and not title:
            flush()
            did = nd
    flush()


def year_from_value(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        y = value.year
        return y if 1990 <= y <= 2100 else None
    if isinstance(value, date):
        y = value.year
        return y if 1990 <= y <= 2100 else None
    return parse_year(value)


def parse_year(value):
    if value is None or value == '':
        return None
    n = parse_num(value)
    if n is not None and 1990 <= n <= 2100:
        return int(n)
    m = re.search(r'(20\d{2})', str(value))
    return int(m.group(1)) if m else None


def pick_common_year(years):
    found = [y for y in years if y]
    if not found:
        return None
    counts = {}
    for y in found:
        counts[y] = counts.get(y, 0) + 1
    return sorted(counts, key=lambda y: (counts[y], y), reverse=True)[0]


def scan_years_rows(rows, max_rows=60):
    found = []
    for row in (rows or [])[:max_rows]:
        for v in row:
            y = parse_year(v)
            if y:
                found.append(y)
    return pick_common_year(found)


def scan_years_sheet(ws, max_row=60, max_col=16):
    found_dt = []
    found_txt = []
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
        for v in row:
            if isinstance(v, (datetime, date)):
                y = year_from_value(v)
                if y:
                    found_dt.append(y)
            else:
                y = parse_year(v)
                if y:
                    found_txt.append(y)
    return pick_common_year(found_dt) or pick_common_year(found_txt)


def scan_first_date(ws, max_row=60, max_col=16):
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
        for v in row:
            if isinstance(v, datetime):
                return v.date().isoformat()
            if isinstance(v, date):
                return v.isoformat()
    return None


def year_from_mtime(path):
    try:
        return datetime.fromtimestamp(os.path.getmtime(path)).year
    except OSError:
        return None


def apply_year_fallback(orgs, path):
    fallback = year_from_mtime(path) or datetime.now().year
    for org in orgs or []:
        if org.get('year') is None:
            org['year'] = fallback


def mean(vals):
    nums = [v for v in vals if v is not None]
    if not nums:
        return None
    return round(sum(nums) / len(nums), 1)


def domain_score(items):
    domain_rows = [e for e in items if not (e.get('sub') or '').strip() and e.get('score') is not None]
    if domain_rows and not any((e.get('sub') or '').strip() for e in items if e.get('score') is not None):
        return mean([e.get('score') for e in domain_rows])
    cap_rows = [e for e in items if (e.get('sub') or '').strip() and e.get('score') is not None]
    if cap_rows:
        return mean([e.get('score') for e in cap_rows])
    return mean([e.get('score') for e in items])


def pillar_score_from_extras(org, did):
    groups = OrderedDict()
    for e in org['extras']:
        if e.get('dirId') != did:
            continue
        name = (e.get('criterion') or e.get('title') or '').strip() or '_'
        groups.setdefault(name, []).append(e)
    if not groups:
        return None
    return mean([domain_score(items) for items in groups.values()])


def year_from_name(name):
    m = re.search(r'(20\d{2})', name or '')
    return int(m.group(1)) if m else None


def name_from_file(filename):
    stem = os.path.splitext(os.path.basename(filename or ''))[0]
    cleaned = re.sub(r'(?i)diaqnostika|reqemsallasma|rəqəmsallaşma|excel|qerar\s*303', ' ', stem)
    cleaned = re.sub(r'[_-]+', ' ', cleaned)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip(' ._-')
    return cleaned or stem or 'Qurum'


def enrich_org(org):
    for e in org['extras']:
        if not e.get('current') and e.get('text') and e.get('text') != e.get('deficiency'):
            e['current'] = e.get('text')
        if not (e.get('deficiency') or '').strip() and e.get('score') is not None:
            e['deficiency'] = infer_gap(
                e.get('score'),
                e.get('current'),
                {},
                e.get('title') or e.get('sub'),
                e.get('dirId'),
            )
        if not e.get('text'):
            e['text'] = e.get('current') or e.get('deficiency') or ''
    for did in DIR_TITLES:
        items = [e for e in org['extras'] if e.get('dirId') == did]
        if not org['dirTexts'].get(did) and len(items) == 1:
            org['dirTexts'][did] = (items[0].get('current') or '').strip()
        if not org['dirGaps'].get(did):
            if items:
                pass
            elif org['dirs'].get(did) is not None and org['dirs'][did] < 100:
                org['dirGaps'][did] = infer_gap(
                    org['dirs'][did],
                    org['dirTexts'].get(did),
                    {},
                    DIR_TITLES[did],
                    did,
                )
    if not org['overallGaps']:
        lows = [e for e in org['extras'] if e.get('score') is not None and e['score'] < 50]
        if lows:
            org['overallGaps'] = '%d altmeyar ilkin və ya idarə olunan səviyyədə qiymətləndirilib.' % len(lows)
        elif org['qrsg'] is not None and org['qrsg'] < 75 and not org['extras']:
            org['overallGaps'] = infer_gap(org['qrsg'], org['overallCurrent'], {}, org['name'], '')


def finalize_orgs(bucket, filename, meta):
    out = []
    fallback = name_from_file(filename)
    meta_score = parse_num(meta.get('score'))
    meta_current = cell_text(meta.get('current'))
    meta_gap = cell_text(meta.get('gap'))
    meta_year = year_from_value(meta.get('year')) or parse_year(meta.get('year')) or year_from_name(filename)
    for key, org in bucket.items():
        if key == '__default__':
            org['name'] = org['name'] or fallback
        if not org['name']:
            org['name'] = fallback
        if org['year'] is None:
            org['year'] = meta_year
        if org['qrsg'] is None and meta_score is not None and (key == '__default__' or len(bucket) == 1):
            org['qrsg'] = meta_score
        if meta_current and not org['overallCurrent'] and len(bucket) == 1:
            org['overallCurrent'] = meta_current
        if meta_gap and not org['overallGaps'] and len(bucket) == 1:
            org['overallGaps'] = meta_gap
        for did in DIR_TITLES:
            if org['dirs'].get(did) is None:
                org['dirs'][did] = pillar_score_from_extras(org, did)
        enrich_org(org)
        dir_vals = [v for v in org['dirs'].values() if v is not None]
        if org['qrsg'] is None and len(dir_vals) >= 3:
            org['qrsg'] = round(sum(dir_vals) / len(dir_vals), 1)
        has = (
            org['qrsg'] is not None
            or any(v is not None for v in org['dirs'].values())
            or org['extras']
            or org['overallCurrent']
            or org['overallGaps']
            or any(org['dirTexts'].values())
            or any(org['dirGaps'].values())
        )
        if has:
            out.append(org)
    return out


def parse_csv_bytes(data, filename):
    text = None
    for enc in ('utf-8-sig', 'utf-8', 'cp1254', 'utf-16'):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = data.decode('utf-8', errors='replace')
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [[cell_text(c) for c in row] for row in reader]
    return parse_rows_as_book(rows, filename, '')


def normalize_rows(rows):
    out = []
    for row in rows or []:
        vals = [cell_text(c) for c in row]
        while vals and not vals[-1]:
            vals.pop()
        out.append(vals)
    while out and not any(out[-1]):
        out.pop()
    return out


def parse_rows_as_book(rows, filename, sheet_dir, meta_extra=None):
    bucket = OrderedDict()
    rows = normalize_rows(rows)
    meta = look_kv(rows)
    if meta_extra:
        for k, v in meta_extra.items():
            meta.setdefault(k, v)
    if not meta.get('year'):
        scanned = scan_years_rows(rows)
        if scanned:
            meta['year'] = scanned
    header_i, cols, marks, levels = find_header(rows)
    if header_i >= 0:
        parse_table(rows, header_i, cols, marks, sheet_dir, meta, bucket, levels)
    else:
        parse_pairs(rows, sheet_dir, meta, bucket)
    return bucket, meta


def parse_diag_excel(path, filename=None):
    filename = filename or os.path.basename(path)
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.csv':
        with open(path, 'rb') as f:
            data = f.read()
        bucket, meta = parse_csv_bytes(data, filename)
        orgs = finalize_orgs(bucket, filename, meta)
        apply_year_fallback(orgs, path)
        return {'orgs': orgs, 'warnings': [] if orgs else ['Cədvəldə diaqnostika sətiri tapılmadı.']}

    wb = require_openpyxl()(path, data_only=True)
    bucket = OrderedDict()
    meta_all = {}
    sheet_year = None
    sheet_date = None
    for ws in wb.worksheets:
        if sheet_year is None:
            sheet_year = scan_years_sheet(ws)
        if sheet_date is None:
            sheet_date = scan_first_date(ws)
        rows = sheet_matrix(ws)
        if not rows:
            continue
        sheet_dir = dir_id_from_title(ws.title)
        part, meta = parse_rows_as_book(rows, filename, sheet_dir)
        for k, v in meta.items():
            meta_all.setdefault(k, v)
        for key, org in part.items():
            if key not in bucket:
                bucket[key] = org
                continue
            dst = bucket[key]
            if org['qrsg'] is not None and dst['qrsg'] is None:
                dst['qrsg'] = org['qrsg']
            if org['year'] and not dst['year']:
                dst['year'] = org['year']
            for did in DIR_TITLES:
                if dst['dirs'].get(did) is None:
                    dst['dirs'][did] = org['dirs'].get(did)
                if org['dirTexts'].get(did):
                    prev = dst['dirTexts'].get(did) or ''
                    dst['dirTexts'][did] = (prev + '\n' + org['dirTexts'][did]).strip() if prev else org['dirTexts'][did]
                if org['dirGaps'].get(did):
                    prev = dst['dirGaps'].get(did) or ''
                    dst['dirGaps'][did] = (prev + '\n' + org['dirGaps'][did]).strip() if prev else org['dirGaps'][did]
            dst['extras'].extend(org['extras'])
            if org['overallCurrent']:
                dst['overallCurrent'] = (dst['overallCurrent'] + '\n' + org['overallCurrent']).strip() if dst['overallCurrent'] else org['overallCurrent']
            if org['overallGaps']:
                dst['overallGaps'] = (dst['overallGaps'] + '\n' + org['overallGaps']).strip() if dst['overallGaps'] else org['overallGaps']
    if sheet_year:
        meta_all.setdefault('year', sheet_year)
    orgs = finalize_orgs(bucket, filename, meta_all)
    apply_year_fallback(orgs, path)
    if sheet_date:
        for org in orgs:
            if not org.get('date'):
                org['date'] = sheet_date
    warnings = []
    if not orgs:
        warnings.append('Excel-də qurum, meyar və ya nəticə cədvəli tapılmadı.')
    return {'orgs': orgs, 'warnings': warnings}
