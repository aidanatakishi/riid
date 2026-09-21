# -*- coding: utf-8 -*-
"""Rəhbər paneli Excel: tab üzrə cədvəl oxuma və şablon."""
from __future__ import annotations

import csv
import io
import os
import re

from diag_excel import cell_text, fold, require_openpyxl

TABS = ('roadmaps', 'services', 'budget', 'arch', 'difficulties')

TAB_LABELS = {
    'roadmaps': 'Yol xəritələri',
    'services': 'Xidmətlərin dizaynı',
    'budget': 'Büdcə və ƏFG',
    'arch': 'Arxitektorlar Qrupu',
    'difficulties': 'Ümumi çətinliklər',
}

SHEET_KIND = (
    ('meta', ('meta', 'umumi', 'basliq', 'yenilenib')),
    ('categories', ('kateqoriya', 'status bolgu', 'yx', 'yol xerite')),
    ('changes', ('deyisiklik', 'kecid', 'change')),
    ('notes', ('qeyd', 'hallar', 'qurum uzre', 'inst')),
    ('services', ('xidmet', 'service', 'reyestr')),
    ('problems', ('problem', 'cetinlik', 'risk', 'blok')),
    ('budget', ('budce', 'budget', 'muraciet')),
    ('efg', ('efg', 'afg', 'kpi kart')),
    ('updates', ('yenilenme', 'update', 'diger')),
    ('work', ('isler', 'gorulen', 'plan', 'arch', 'arxitektur')),
    ('difficulties', ('umumi cetinlik', 'cetinlikler', 'riskler')),
)

COLS = {
    'category': ('kateqoriya', 'status', 'qrup'),
    'title': ('basliq', 'ad', 'title'),
    'desc': ('tesvir', 'aciqlama', 'alt', 'subtitle'),
    'prev': ('evvelki', 'onceki', 'previous'),
    'current': ('cari', 'indiki', 'current', 'say', 'sayi'),
    'qurum': ('qurum', 'teskilat', 'chip', 'qurumlar'),
    'new': ('yeni', 'new'),
    'from_to': ('kecid', 'from', 'to', 'deyisiklik'),
    'text': ('metn', 'qeyd', 'note', 'text', 'mezmun'),
    'tag': ('tag', 'etiket', 'kim', 'who', 'bolme'),
    'badge': ('badge', 'qisa', 'abbrev'),
    'status': ('status', 'pill', 'veziyyet'),
    'headsup': ('headsup', 'heads up', 'rehberlik', 'teleb'),
    'update': ('yenilenme', 'update'),
    'list1_title': ('siyahi 1 basliq', 'siyahi basligi', 'siyahi 1 adi'),
    'list1': ('siyahi 1', 'siyahi', 'siyahı 1'),
    'list2_title': ('siyahi 2 basliq', 'siyahi 2 adi'),
    'list2': ('siyahi 2',),
    'service': ('xidmet', 'service'),
    'old_status': ('evvelki status', 'onceki status'),
    'business': ('biznes', 'business'),
    'kind': ('nov', 'bolme', 'kind', 'tip'),
    'who': ('who', 'kim', 'qurum', 'layihe'),
    'value': ('deyer', 'reqem', 'value', 'n'),
    'key': ('acar', 'gosteric', 'key', 'ad'),
}


def _has_needle(folded, needles):
    return any(n in folded for n in needles)


def classify_sheet(name):
    f = fold(name)
    for kind, needles in SHEET_KIND:
        if _has_needle(f, needles):
            return kind
    return ''


def split_list(raw):
    text = cell_text(raw)
    if not text:
        return []
    parts = re.split(r'[\n;|]+', text)
    if len(parts) == 1 and ',' in text and text.count(',') >= 1:
        parts = text.split(',')
    out = []
    seen = set()
    for part in parts:
        item = part.strip().strip('·').strip()
        if not item or item in seen:
            continue
        seen.add(item)
        out.append(item)
    return out


def num_or_none(raw):
    text = cell_text(raw).replace(' ', '').replace('%', '').replace(',', '.')
    if not text:
        return None
    try:
        n = float(text)
    except ValueError:
        return None
    if n == int(n):
        return int(n)
    return n


def truthy(raw):
    f = fold(raw)
    return f in ('1', 'true', 'yes', 'beli', 'he', 'yeni', 'x', 'v', 'ok')


def map_headers(row):
    mapping = {}
    for i, cell in enumerate(row or []):
        f = fold(cell)
        if not f:
            continue
        for key, needles in COLS.items():
            if key in mapping:
                continue
            if _has_needle(f, needles) or f == fold(needles[0]):
                mapping[key] = i
                break
    return mapping


def row_get(row, mapping, key, default=''):
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return default
    return cell_text(row[idx])


def load_tables(path):
    ext = os.path.splitext(path)[1].lower()
    tables = {}
    if ext == '.csv':
        with open(path, 'r', encoding='utf-8-sig', newline='') as handle:
            rows = [[cell_text(c) for c in row] for row in csv.reader(handle)]
        tables['Sheet1'] = [r for r in rows if any(r)]
        return tables
    load_workbook = require_openpyxl()
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            rows = []
            for raw in ws.iter_rows(values_only=True):
                cells = [cell_text(c) for c in raw]
                if any(cells):
                    rows.append(cells)
            tables[ws.title or 'Sheet'] = rows
    finally:
        wb.close()
    return tables


def rows_as_dicts(rows):
    if not rows:
        return []
    header_i = 0
    mapping = {}
    for i, row in enumerate(rows[:8]):
        found = map_headers(row)
        if len(found) >= 2 or (len(found) == 1 and i == 0):
            mapping = found
            header_i = i
            break
    if not mapping:
        return []
    out = []
    for row in rows[header_i + 1:]:
        if not any(row):
            continue
        item = {}
        for key, idx in mapping.items():
            item[key] = cell_text(row[idx]) if idx < len(row) else ''
        if any(item.values()):
            out.append(item)
    return out


def parse_meta(rows):
    meta = {}
    for row in rows or []:
        if len(row) < 2:
            continue
        key = fold(row[0])
        val = cell_text(row[1])
        if not key or not val:
            continue
        if _has_needle(key, ('yenilen', 'tarix', 'date')):
            meta['updated'] = val
        elif _has_needle(key, ('altbasliq', 'sub', 'tesvir')):
            meta['sub'] = val
        elif _has_needle(key, ('yx say', 'cem', 'total', 'seal')):
            meta['total'] = num_or_none(val)
        elif _has_needle(key, ('eyebrow', 'basliq xett')):
            meta['eyebrow'] = val
    dicts = rows_as_dicts(rows)
    for item in dicts:
        key = fold(item.get('key') or item.get('title') or '')
        val = item.get('value') or item.get('text') or item.get('current') or ''
        if _has_needle(key, ('yenilen', 'tarix')):
            meta['updated'] = val
        elif _has_needle(key, ('altbasliq', 'sub')):
            meta['sub'] = val
        elif _has_needle(key, ('yx', 'cem', 'total')):
            meta['total'] = num_or_none(val)
    return meta


def parse_categories(rows):
    items = []
    for row in rows_as_dicts(rows):
        title = row.get('category') or row.get('title') or ''
        qurums = split_list(row.get('qurum'))
        if not title and not qurums:
            continue
        current = len(qurums) if qurums else num_or_none(row.get('current'))
        prev = num_or_none(row.get('prev'))
        items.append({
            'title': title,
            'desc': row.get('desc') or '',
            'prev': prev,
            'current': current,
            'qurums': [{'name': name, 'isNew': 'yeni' in fold(name)} for name in qurums],
        })
    if items:
        return items
    grouped = {}
    order = []
    for row in rows_as_dicts(rows):
        cat = row.get('category') or ''
        name = row.get('qurum') or row.get('title') or ''
        if not cat or not name:
            continue
        if cat not in grouped:
            grouped[cat] = {
                'title': cat,
                'desc': row.get('desc') or '',
                'prev': num_or_none(row.get('prev')),
                'qurums': []
            }
            order.append(cat)
        grouped[cat]['qurums'].append({'name': name, 'isNew': truthy(row.get('new')) or 'yeni' in fold(name)})
    out = []
    for cat in order:
        rec = grouped[cat]
        rec['current'] = len(rec['qurums'])
        out.append(rec)
    return out


def parse_changes(rows):
    out = []
    for row in rows_as_dicts(rows):
        tag = row.get('from_to') or row.get('tag') or row.get('title') or ''
        text = row.get('text') or row.get('desc') or ''
        if tag or text:
            out.append({'tag': tag, 'text': text})
    return out


def parse_notes(rows):
    out = []
    for row in rows_as_dicts(rows):
        name = row.get('title') or row.get('qurum') or ''
        if not name and not row.get('text'):
            continue
        lists = []
        if row.get('list1'):
            lists.append({'title': row.get('list1_title') or '', 'items': split_list(row.get('list1'))})
        if row.get('list2'):
            lists.append({'title': row.get('list2_title') or '', 'items': split_list(row.get('list2'))})
        out.append({
            'badge': row.get('badge') or '',
            'name': name,
            'status': row.get('status') or '',
            'text': row.get('text') or row.get('desc') or '',
            'headsup': row.get('headsup') or '',
            'update': row.get('update') or '',
            'lists': lists,
        })
    return out


def parse_services(rows):
    out = []
    for row in rows_as_dicts(rows):
        org = row.get('qurum') or ''
        service = row.get('service') or row.get('title') or ''
        if not org and not service:
            continue
        status = row.get('status') or ''
        biz = truthy(row.get('business')) or 'biznes' in fold(service)
        out.append({
            'org': org,
            'service': service,
            'status': status,
            'oldStatus': row.get('old_status') or '',
            'note': row.get('text') or row.get('desc') or '',
            'business': biz,
        })
    return out


def parse_problems(rows):
    out = []
    for row in rows_as_dicts(rows):
        tag = row.get('tag') or row.get('title') or row.get('qurum') or ''
        text = row.get('text') or row.get('desc') or ''
        deadline = ''
        kind = fold(row.get('kind') or '')
        if _has_needle(kind, ('deadline', 'tecili', 'son tarix')):
            deadline = text
            text = row.get('desc') or text
        if tag or text:
            out.append({'tag': tag, 'text': text, 'deadline': deadline, 'kind': row.get('kind') or ''})
    return out


def parse_stats_pairs(rows):
    items = []
    for row in rows_as_dicts(rows):
        label = row.get('key') or row.get('title') or row.get('category') or row.get('desc') or ''
        value = num_or_none(row.get('value') if row.get('value') != '' else row.get('current'))
        if not label and value is None:
            continue
        items.append({'label': label, 'value': value})
    if items:
        return items
    if rows and len(rows[0]) >= 2 and not map_headers(rows[0]):
        for row in rows:
            if len(row) < 2:
                continue
            label = cell_text(row[0])
            value = num_or_none(row[1])
            if label:
                items.append({'label': label, 'value': value})
    return items


def parse_work(rows):
    out = []
    for row in rows_as_dicts(rows):
        kind = row.get('kind') or row.get('tag') or ''
        who = row.get('who') or row.get('qurum') or ''
        title = row.get('title') or ''
        text = row.get('text') or row.get('desc') or ''
        if not (kind or who or title or text):
            continue
        folded = fold(kind)
        section = 'done'
        if _has_needle(folded, ('davam', 'yarimciq', '2')):
            section = 'ongoing'
        elif _has_needle(folded, ('problem', 'blok', '3')):
            section = 'block'
        elif _has_needle(folded, ('plan', 'hefte', '4')):
            section = 'plan'
        elif _has_needle(folded, ('gorulen', 'tamam', '1')):
            section = 'done'
        out.append({'section': section, 'who': who, 'title': title, 'text': text})
    return out


def parse_difficulties(rows):
    out = []
    for row in rows_as_dicts(rows):
        text = row.get('text') or row.get('desc') or row.get('title') or ''
        if text:
            out.append({'text': text})
    if out:
        return out
    for row in rows or []:
        text = cell_text(row[0] if row else '')
        f = fold(text)
        if text and not _has_needle(f, ('metn', 'cetinlik', 'qeyd')):
            out.append({'text': text})
    return out


def _pick_sheet(tables, kinds, fallback_first=False):
    named = []
    for name, rows in tables.items():
        kind = classify_sheet(name)
        if kind in kinds:
            named.append(rows)
    if named:
        return named[0]
    if fallback_first and tables:
        return list(tables.values())[0]
    return []


def parse_tab(path, tab):
    tab = str(tab or '').strip().lower()
    if tab not in TABS:
        raise ValueError('Naməlum tab')
    tables = load_tables(path)
    if not tables:
        raise ValueError('Excel boşdur')
    meta = parse_meta(_pick_sheet(tables, ('meta',)))
    if tab == 'roadmaps':
        cats = parse_categories(_pick_sheet(tables, ('categories',), True))
        data = {
            'meta': meta,
            'categories': cats,
            'changes': parse_changes(_pick_sheet(tables, ('changes',))),
            'notes': parse_notes(_pick_sheet(tables, ('notes',))),
        }
        if not cats and not data['changes'] and not data['notes']:
            raise ValueError('Yol xəritələri cədvəli tapılmadı. Şablondakı sütun adlarını saxlayın.')
        return data
    if tab == 'services':
        services = parse_services(_pick_sheet(tables, ('services',), True))
        data = {
            'meta': meta,
            'services': services,
            'problems': parse_problems(_pick_sheet(tables, ('problems', 'difficulties'))),
        }
        if not services:
            raise ValueError('Xidmət cədvəli tapılmadı. Qurum / Xidmət / Status sütunları lazımdır.')
        return data
    if tab == 'budget':
        data = {
            'meta': meta,
            'budget': parse_stats_pairs(_pick_sheet(tables, ('budget',), True)),
            'efg': parse_stats_pairs(_pick_sheet(tables, ('efg',))),
            'problems': parse_problems(_pick_sheet(tables, ('problems', 'difficulties'))),
            'updates': parse_changes(_pick_sheet(tables, ('updates', 'changes'))),
        }
        if not data['budget'] and not data['efg'] and not data['problems']:
            raise ValueError('Büdcə / ƏFG cədvəli tapılmadı.')
        return data
    if tab == 'arch':
        work = parse_work(_pick_sheet(tables, ('work',), True))
        data = {
            'meta': meta,
            'work': work,
            'problems': parse_problems(_pick_sheet(tables, ('problems', 'difficulties'))),
        }
        if not work and not data['problems']:
            raise ValueError('Arxitektorlar cədvəli tapılmadı. Bölmə / Kim / Mətn sütunları lazımdır.')
        return data
    diffs = parse_difficulties(_pick_sheet(tables, ('difficulties', 'problems'), True))
    if not diffs:
        raise ValueError('Çətinliklər cədvəli tapılmadı. Mətn sütunu lazımdır.')
    return {'meta': meta, 'items': diffs}


def _write_sheet(wb, title, headers, sample=None):
    ws = wb.create_sheet(title)
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    if sample:
        for i, val in enumerate(sample, 1):
            ws.cell(2, i, val)
    return ws


def build_template(tab):
    from openpyxl import Workbook
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    tab = str(tab or '').strip().lower()
    if tab == 'roadmaps':
        _write_sheet(wb, 'Meta', ['Açar', 'Dəyər'])
        _write_sheet(wb, 'Kateqoriyalar', ['Kateqoriya', 'Təsvir', 'Əvvəlki', 'Cari', 'Qurumlar'])
        _write_sheet(wb, 'Dəyişikliklər', ['Keçid', 'Mətn'])
        _write_sheet(wb, 'Qeydlər', ['Badge', 'Qurum', 'Status', 'Mətn', 'Heads-up', 'Yenilənmə', 'Siyahı 1 başlığı', 'Siyahı 1', 'Siyahı 2 başlığı', 'Siyahı 2'])
    elif tab == 'services':
        _write_sheet(wb, 'Meta', ['Açar', 'Dəyər'])
        _write_sheet(wb, 'Xidmətlər', ['Qurum', 'Xidmət', 'Status', 'Əvvəlki status', 'Qeyd', 'Biznes'])
        _write_sheet(wb, 'Problemlər', ['Tag', 'Mətn'])
    elif tab == 'budget':
        _write_sheet(wb, 'Meta', ['Açar', 'Dəyər'])
        _write_sheet(wb, 'Büdcə', ['Göstərici', 'Dəyər'])
        _write_sheet(wb, 'ƏFG', ['Göstərici', 'Dəyər'])
        _write_sheet(wb, 'Problemlər', ['Tag', 'Mətn', 'Növ'])
        _write_sheet(wb, 'Yenilənmə', ['Keçid', 'Mətn'])
    elif tab == 'arch':
        _write_sheet(wb, 'Meta', ['Açar', 'Dəyər'])
        _write_sheet(wb, 'İşlər', ['Bölmə', 'Kim', 'Başlıq', 'Mətn'])
        _write_sheet(wb, 'Problemlər', ['Tag', 'Mətn'])
    elif tab == 'difficulties':
        _write_sheet(wb, 'Çətinliklər', ['Mətn'])
    else:
        raise ValueError('Naməlum tab')
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def template_filename(tab):
    label = TAB_LABELS.get(tab, tab)
    safe = re.sub(r'[^\w\-]+', '_', fold(label)) or tab
    return 'rehber_' + safe + '_sablon.xlsx'
